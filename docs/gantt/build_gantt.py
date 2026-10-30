"""Convert the supplied month-level VED schedule into native editable draw.io cells."""
import hashlib
import html
import json
import re
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
SOURCE = Path('C:/Users/kupal/Downloads/VEDGantt-Chart.xlsx')
REFERENCE = Path('C:/Users/kupal/Downloads/SmartSpace AR_ A Personalized, Context-Aware Augmented Reality Platform for Interactive Interior Design and Space Optimization Manuscript.pdf')
SOURCE_HASH = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
REFERENCE_HASH = hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
BOOK = openpyxl.load_workbook(SOURCE, data_only=True)
SHEET = BOOK['Thesis Gantt']
MONTHS = ['JANUARY', 'FEBRUARY', 'MARCH', 'APRIL', 'MAY', 'JUNE', 'JULY', 'AUGUST', 'SEPTEMBER', 'OCTOBER']
SOURCE_COLORS = {'FF548235': 'Completed', 'FFED7D31': 'Planning', 'FF1F3864': 'Phase schedule'}
COLORS = {'Completed': '#93DA9A', 'Planning': '#FFAD80', 'Phase schedule': '#1F3864'}
FONT_CACHE = {}


def font(size, bold=False):
    key = (size, bold)
    if key not in FONT_CACHE:
        FONT_CACHE[key] = ImageFont.truetype('C:/Windows/Fonts/' + ('arialbd.ttf' if bold else 'arial.ttf'), size)
    return FONT_CACHE[key]


def clean(value):
    # The workbook contains replacement characters in dash positions.
    return str(value).replace('\ufffd', '-').replace('\u2013', '-').replace('\u2014', '-')


def wrap(value, width, size, bold=False):
    lines = []
    for paragraph in value.split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = (line + ' ' + word).strip()
            if line and font(size, bold).getlength(candidate) > width:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    assert all(font(size, bold).getlength(line) <= width for line in lines), value
    return lines


def source_fill(cell):
    if cell.fill.patternType == 'solid' and cell.fill.fgColor.type == 'rgb':
        return SOURCE_COLORS.get(cell.fill.fgColor.rgb)
    return None


groups = []
for r in range(4, 43):
    name = SHEET.cell(r, 1).value
    if name is None:
        continue
    status = SHEET.cell(r, 2).value
    schedule = {m: source_fill(SHEET.cell(r, 4 + m)) for m in range(10)}
    active = {m: kind for m, kind in schedule.items() if kind}
    if status is None:
        label = clean(name)
        label = re.sub(r'\s*\([^)]*2026\)\s*$', '', label)
        label = re.sub(r'^PHASE (\d+)\s*-\s*', r'PHASE \1: ', label)
        label = label.replace('DOCUMENTATION - ', 'DOCUMENTATION: ')
        groups.append({'row': r, 'raw_name': name, 'label': label, 'schedule': active, 'tasks': []})
    else:
        assert status in {'Completed', 'Planning'}
        assert active and all(kind == status for kind in active.values()), (r, status, active)
        groups[-1]['tasks'].append({'row': r, 'raw_name': name, 'name': clean(name), 'status': status,
                                    'raw_place': SHEET.cell(r, 3).value, 'place': clean(SHEET.cell(r, 3).value),
                                    'schedule': active})

assert len(groups) == 4
all_tasks = [task for group in groups for task in group['tasks']]
assert len(all_tasks) == 35
assert not any(c.data_type == 'f' for row in openpyxl.load_workbook(SOURCE, data_only=False).active for c in row)

# Keep every source task once. This schedule has no task that crosses the two panels.
PANELS = []
for start, end, title in [(0, 6, 'Gantt Chart - 2026 (January - June)'), (6, 10, 'Gantt Chart - 2026 (July - October)')]:
    selected = []
    for group in groups:
        tasks = [task for task in group['tasks'] if any(start <= m < end for m in task['schedule'])]
        if tasks:
            selected.append({'group': group, 'tasks': tasks})
    PANELS.append({'start': start, 'end': end, 'title': title, 'groups': selected})
displayed = [task['row'] for panel in PANELS for group in panel['groups'] for task in group['tasks']]
assert len(displayed) == len(set(displayed)) == len(all_tasks)
assert set(displayed) == {task['row'] for task in all_tasks}

W = 1470
MARGIN = 40
LEFT = 315
RIGHT = 259
TIMELINE = W - MARGIN * 2 - LEFT - RIGHT
MONTH_Y = 36
WEEK_Y = 30
PHASE_H = 42


def task_height(task):
    name_h = len(wrap(task['name'], LEFT - 20, 13, True)) * 17 + 14
    place_h = len(wrap(task['place'], RIGHT - 16, 11)) * 15 + 35
    return max(48, name_h, place_h)


for panel in PANELS:
    panel['table_height'] = MONTH_Y + WEEK_Y + sum(PHASE_H + sum(task_height(task) for task in g['tasks']) for g in panel['groups'])
H = 96 + sum(58 + panel['table_height'] + 45 for panel in PANELS) + 123
MX = ET.Element('mxfile', {'host': 'app.diagrams.net', 'type': 'device', 'version': '24.7.17'})
DIAGRAM = ET.SubElement(MX, 'diagram', {'id': 'ved-gantt-2026', 'name': 'VED Gantt Chart - 2026'})
MODEL = ET.SubElement(DIAGRAM, 'mxGraphModel', {'dx': str(W), 'dy': str(H), 'grid': '1', 'gridSize': '10', 'guides': '1', 'connect': '1', 'arrows': '1', 'page': '1', 'pageScale': '1', 'pageWidth': str(W), 'pageHeight': str(H), 'background': '#FFFFFF'})
ROOT = ET.SubElement(MODEL, 'root')
ET.SubElement(ROOT, 'mxCell', {'id': '0'})
ET.SubElement(ROOT, 'mxCell', {'id': '1', 'parent': '0'})
IMAGE = Image.new('RGB', (W, H), 'white')
DRAW = ImageDraw.Draw(IMAGE)
SVG = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}"><rect width="{W}" height="{H}" fill="white"/>']
SEQ = 0
SHAPES = []
CELL_MODEL = []


def group(cid, x, y, w, h, parent='1', px=0, py=0):
    c = ET.SubElement(ROOT, 'mxCell', {'id': cid, 'value': '', 'vertex': '1', 'parent': parent, 'style': 'group;'})
    ET.SubElement(c, 'mxGeometry', {'x': str(x-px), 'y': str(y-py), 'width': str(w), 'height': str(h), 'as': 'geometry'})
    return cid


def rect(x, y, w, h, fill='white', stroke='#555555', parent='1', px=0, py=0, cid=None, source=None):
    global SEQ
    SEQ += 1
    cid = cid or f'r{SEQ}'
    assert 0 <= x <= W-w and 0 <= y <= H-h
    attrs = {'id': cid, 'value': '', 'vertex': '1', 'parent': parent,
             'style': f'rounded=0;html=0;fillColor={fill};strokeColor={stroke};strokeWidth=0.8;'}
    if source:
        attrs['sourceCell'] = source
    c = ET.SubElement(ROOT, 'mxCell', attrs)
    ET.SubElement(c, 'mxGeometry', {'x': str(x-px), 'y': str(y-py), 'width': str(w), 'height': str(h), 'as': 'geometry'})
    SVG.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="{stroke}" stroke-width="0.8"/>')
    DRAW.rectangle((x, y, x+w, y+h), fill=fill, outline=stroke, width=1)
    SHAPES.append((cid, x, y, w, h))
    return cid


def text(value, x, y, w, h, size=13, bold=False, align='center', parent='1', px=0, py=0, source=None):
    global SEQ
    SEQ += 1
    lines = wrap(value, w-10, size, bold)
    lh = size + 4
    assert len(lines)*lh <= h, ('Text overflow', value, h, lines)
    assert 0 <= x <= W-w and 0 <= y <= H-h
    attrs = {'id': f't{SEQ}', 'value': '\n'.join(lines), 'vertex': '1', 'parent': parent,
             'style': f'text;html=0;whiteSpace=wrap;fillColor=none;strokeColor=none;fontFamily=Arial;fontSize={size};fontStyle={int(bold)};fontColor=#000000;align={align};verticalAlign=middle;spacing=5;'}
    if source:
        attrs['sourceCell'] = source
    c = ET.SubElement(ROOT, 'mxCell', attrs)
    ET.SubElement(c, 'mxGeometry', {'x': str(x-px), 'y': str(y-py), 'width': str(w), 'height': str(h), 'as': 'geometry'})
    ty = y+(h-len(lines)*lh)/2
    for line in lines:
        length = font(size, bold).getlength(line)
        tx = x+5 if align == 'left' else x+(w-length)/2
        SVG.append(f'<text x="{tx}" y="{ty+size}" font-family="Arial,sans-serif" font-size="{size}" font-weight="{700 if bold else 400}">{html.escape(line)}</text>')
        DRAW.text((tx, ty), line, font=font(size, bold), fill='black')
        ty += lh


text('Appendix Table. Gantt Chart', MARGIN, 20, W-2*MARGIN, 35, 22, True, 'left')
y = 96
for pi, panel in enumerate(PANELS, 1):
    title_width = 530
    rect((W-title_width)/2, y, title_width, 40, stroke='#888888')
    text(panel['title'], (W-title_width)/2, y, title_width, 40, 20, True)
    y += 58
    table_y = y
    table_group = group(f'panel-{pi}', MARGIN, y, W-2*MARGIN, panel['table_height'])
    x0 = MARGIN+LEFT
    last = x0+TIMELINE
    nm = panel['end']-panel['start']
    cw = TIMELINE/(nm*4)
    rect(MARGIN, y, LEFT, MONTH_Y+WEEK_Y, parent=table_group, px=MARGIN, py=table_y)
    text('RESEARCH ACTIVITIES', MARGIN, y, LEFT, MONTH_Y+WEEK_Y, 13, True, parent=table_group, px=MARGIN, py=table_y)
    rect(last, y, RIGHT, MONTH_Y+WEEK_Y, parent=table_group, px=MARGIN, py=table_y)
    text('STATUS / DESCRIPTION', last, y, RIGHT, MONTH_Y+WEEK_Y, 13, True, parent=table_group, px=MARGIN, py=table_y)
    for mi, m in enumerate(range(panel['start'], panel['end'])):
        mx = x0+mi*cw*4
        rect(mx, y, cw*4, MONTH_Y, parent=table_group, px=MARGIN, py=table_y)
        text(MONTHS[m], mx, y, cw*4, MONTH_Y, 14, True, parent=table_group, px=MARGIN, py=table_y)
        for week in range(4):
            xx = mx+week*cw
            rect(xx, y+MONTH_Y, cw, WEEK_Y, parent=table_group, px=MARGIN, py=table_y)
            text(f'W{week+1}', xx, y+MONTH_Y, cw, WEEK_Y, 12, parent=table_group, px=MARGIN, py=table_y)
    y += MONTH_Y+WEEK_Y
    for g in panel['groups']:
        phase = g['group']
        row_group = group(f'phase-{pi}-{phase["row"]}', MARGIN, y, W-2*MARGIN, PHASE_H, table_group, MARGIN, table_y)
        rect(MARGIN, y, LEFT, PHASE_H, '#F2F2F2', parent=row_group, px=MARGIN, py=y)
        text(phase['label'], MARGIN+5, y, LEFT-10, PHASE_H, 12, True, 'left', row_group, MARGIN, y, f'A{phase["row"]}')
        rect(last, y, RIGHT, PHASE_H, '#F2F2F2', parent=row_group, px=MARGIN, py=y)
        text('PHASE SCHEDULE', last, y, RIGHT, PHASE_H, 12, True, parent=row_group, px=MARGIN, py=y)
        for mi, m in enumerate(range(panel['start'], panel['end'])):
            for week in range(4):
                filled = m in phase['schedule']
                rect(x0+(mi*4+week)*cw, y, cw, PHASE_H, COLORS['Phase schedule'] if filled else '#F2F2F2',
                     parent=row_group, px=MARGIN, py=y, cid=f'phase-{pi}-{phase["row"]}-m{m}-w{week}', source=SHEET.cell(phase['row'], 4+m).coordinate)
        y += PHASE_H
        for task in g['tasks']:
            rh = task_height(task)
            row_group = group(f'task-{task["row"]}', MARGIN, y, W-2*MARGIN, rh, table_group, MARGIN, table_y)
            rect(MARGIN, y, LEFT, rh, parent=row_group, px=MARGIN, py=y)
            text(task['name'], MARGIN+5, y, LEFT-10, rh, 13, True, 'left', row_group, MARGIN, y, f'A{task["row"]}')
            rect(last, y, RIGHT, rh, parent=row_group, px=MARGIN, py=y)
            text(task['status'].upper(), last, y+4, RIGHT, 17, 11, True, parent=row_group, px=MARGIN, py=y, source=f'B{task["row"]}')
            text(task['place'], last+3, y+24, RIGHT-6, rh-28, 11, parent=row_group, px=MARGIN, py=y, source=f'C{task["row"]}')
            for mi, m in enumerate(range(panel['start'], panel['end'])):
                for week in range(4):
                    kind = task['schedule'].get(m)
                    cid = rect(x0+(mi*4+week)*cw, y, cw, rh, COLORS[kind] if kind else '#FFFFFF',
                               parent=row_group, px=MARGIN, py=y, cid=f'task-{task["row"]}-m{m}-w{week}', source=SHEET.cell(task['row'], 4+m).coordinate)
                    CELL_MODEL.append({'id': cid, 'source_row': task['row'], 'month': m, 'week': week+1, 'status': kind})
            y += rh
    assert y-table_y == panel['table_height']
    # Compact horizontal legend below each table, like the manuscript reference.
    ly = y+9
    labels = [('No activity', '#FFFFFF'), ('Completed', COLORS['Completed']), ('Planning', COLORS['Planning']), ('Phase schedule', COLORS['Phase schedule'])]
    for i, (label, color) in enumerate(labels):
        xx = MARGIN+LEFT+20+i*250
        rect(xx, ly+4, 23, 12, color, '#777777')
        text(label, xx+30, ly, 180, 21, 12, False, 'left')
    y += 45

text('Gantt Chart of VED Electrical Services', MARGIN, y+8, W-2*MARGIN, 35, 21, False)
text('AI-Driven Floor Plan Analysis and 3D Visualization System for Automated Layout and Cost Estimation', MARGIN, y+44, W-2*MARGIN, 28, 16)
text('Month-level schedule: each active month is shown across W1-W4. Specific weekly dates were not supplied.', MARGIN, y+83, W-2*MARGIN, 25, 13)
SVG.append('</svg>')
ET.indent(MX)
FILE = OUT/'ved-gantt-chart-2026-10-14.drawio'
ET.ElementTree(MX).write(FILE, encoding='utf-8', xml_declaration=True)
IMAGE.save(OUT/'ved-gantt-chart.png')
(OUT/'ved-gantt-chart.svg').write_text(''.join(SVG), encoding='utf-8')
(OUT/'preview.html').write_text('<!doctype html><html><meta charset="utf-8"><title>VED Gantt chart</title><style>body{margin:0;background:#ddd}svg{display:block;width:100%;max-width:1470px;margin:auto;background:white}</style>'+''.join(SVG)+'</html>', encoding='utf-8')
(OUT/'source-schedule.json').write_text(json.dumps({'sheet': SHEET.title, 'groups': groups, 'panels': [{'title': p['title'], 'task_rows': [t['row'] for g in p['groups'] for t in g['tasks']]} for p in PANELS]}, indent=2), encoding='utf-8')

parsed = ET.parse(FILE)
assert len(parsed.findall('diagram')) == 1
cells = parsed.findall('.//mxCell')
ids = {c.attrib['id'] for c in cells}
assert len(ids) == len(cells)
assert all(c.attrib.get('parent', '0') in ids for c in cells)
assert not any('image' in c.attrib.get('style', '') for c in cells)
by_id = {c.attrib['id']: c for c in cells}
for task in all_tasks:
    for column, expected in [('A', task['name']), ('B', task['status'].upper()), ('C', task['place'])]:
        matches = [c for c in cells if c.attrib.get('sourceCell') == f'{column}{task["row"]}']
        assert len(matches) == 1
        assert ' '.join(matches[0].attrib['value'].split()) == ' '.join(expected.split())
source_months = sum(len(task['schedule']) for task in all_tasks)
assert sum(c['status'] is not None for c in CELL_MODEL) == source_months*4
for task in all_tasks:
    for m in range(10):
        mapped = [c for c in CELL_MODEL if c['source_row'] == task['row'] and c['month'] == m]
        if task['schedule'].get(m):
            assert len(mapped) == 4 and all(c['status'] == task['status'] for c in mapped)
        else:
            assert all(c['status'] is None for c in mapped)
for model in CELL_MODEL:
    expected = COLORS[model['status']] if model['status'] else '#FFFFFF'
    assert f'fillColor={expected};' in by_id[model['id']].attrib['style']
    assert by_id[model['id']].attrib['sourceCell'] == SHEET.cell(model['source_row'], 4+model['month']).coordinate
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == SOURCE_HASH
assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == REFERENCE_HASH
report = {'pages': 1, 'stacked_tables': 2, 'tasks': len(all_tasks), 'status_counts': dict(Counter(t['status'] for t in all_tasks)),
          'period': 'January-October 2026', 'months_from_source': 10, 'scheduled_task_months': source_months,
          'colored_task_week_cells': source_months*4, 'monthly_source_mapping': 'PASS', 'every_task_once': 'PASS',
          'task_status_and_place_preservation': 'PASS', 'native_editable_xml': 'PASS', 'text_fit': 'PASS', 'source_files_unchanged': 'PASS',
          'source_workbook_sha256': SOURCE_HASH, 'source_reference_sha256': REFERENCE_HASH,
          'source_note': 'Phase 2 heading says Aug-Sep, but its source phase fill includes Aug-Oct. Chart retains source filled months and omits date ranges from phase labels.',
          'weekly_detail': 'Not supplied: W1-W4 are visual subdivisions of month-level source data.',
          'visual_review': 'PENDING', 'editor_import': 'NOT TESTED', 'downloads_copy': 'PENDING'}
(OUT/'verification.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
print(json.dumps(report, indent=2))
